"""Private uploads, staged native-vector indices, and live permission filtering."""
from __future__ import annotations

import asyncio
import io
import json
import math
import zipfile
from contextlib import closing
from pathlib import Path

import sqlite_vec
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pydantic import BaseModel, Field

from . import db, storage, tasks
from .auth import require_owner
from .config import settings
from .models_routes import client_for, config_for, config_at

router = APIRouter(prefix="/api/v1/admin/knowledge", tags=["knowledge"], dependencies=[Depends(require_owner)])


def init_schema():
    with closing(db.connect()) as conn:
        conn.executescript("""
        CREATE TABLE IF NOT EXISTS knowledge_sources (
          id TEXT PRIMARY KEY, object_id TEXT, content_id TEXT UNIQUE, filename TEXT NOT NULL,
          visibility TEXT NOT NULL DEFAULT 'private', rag_enabled INTEGER NOT NULL DEFAULT 0,
          revision INTEGER NOT NULL DEFAULT 1, active_version TEXT, status TEXT NOT NULL,
          text TEXT, error TEXT, task_id TEXT, created_at TEXT NOT NULL, updated_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS knowledge_versions (
          id TEXT PRIMARY KEY, source_id TEXT NOT NULL, source_revision INTEGER NOT NULL,
          model_id TEXT NOT NULL, model_version INTEGER NOT NULL, dimensions INTEGER,
          status TEXT NOT NULL, created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS knowledge_chunks (
          id TEXT PRIMARY KEY, source_id TEXT NOT NULL, version_id TEXT NOT NULL,
          position INTEGER NOT NULL, text TEXT NOT NULL, embedding BLOB NOT NULL
        );
        CREATE INDEX IF NOT EXISTS knowledge_chunks_version ON knowledge_chunks(version_id);
        """)
    tasks.register("knowledge.index", index_source)


def extract_text(data: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in (".md", ".txt"):
        try:
            result = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ValueError("文本文件必须采用 UTF-8 编码") from None
    elif suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError("暂不支持加密 PDF")
        if len(reader.pages) > 1000:
            raise ValueError("PDF 页数超过 1000 页")
        result = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == ".docx":
        from docx import Document
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("DOCX 解压大小超过 100MiB")
        document = Document(io.BytesIO(data))
        result = "\n".join([p.text for p in document.paragraphs] + [" | ".join(c.text for c in row.cells) for table in document.tables for row in table.rows])
    else:
        raise ValueError("仅支持 md、txt、pdf、docx")
    result = result.replace("\x00", "").strip()
    if not result:
        raise ValueError("文件没有可提取文本；扫描 PDF 请先完成 OCR")
    if len(result) > 2_000_000:
        raise ValueError("提取文本超过 200 万字符，请拆分上传")
    return result


def source_record(source_id: str):
    with closing(db.connect()) as conn:
        row = conn.execute("SELECT * FROM knowledge_sources WHERE id=?", (source_id,)).fetchone()
    if not row:
        raise HTTPException(404, "知识来源不存在")
    return dict(row)


def live_allowed(source: dict, owner: bool = False) -> bool:
    if owner:
        return True
    if source.get("content_id"):
        with closing(db.connect()) as conn:
            row = conn.execute("SELECT draft_json,published_json FROM contents WHERE id=?", (source["content_id"],)).fetchone()
        if not row or not row["published_json"]:
            return False
        draft, published = json.loads(row["draft_json"]), json.loads(row["published_json"])
        return draft.get("visibility") == "public" and bool(draft.get("rag_enabled")) and published.get("visibility") == "public" and bool(published.get("rag_enabled"))
    return source["visibility"] == "public" and bool(source["rag_enabled"])


async def embed_texts(config: dict, texts: list[str], request_id: str | None = None, attempt_id: str | None = None) -> list[list[float]]:
    call_id = db.uid()
    with db.transaction() as conn:
        conn.execute("INSERT INTO model_calls(id,request_id,attempt_id,purpose,model_id,model_version,status,created_at) VALUES (?,?,?,?,?,?,?,?)", (call_id, request_id, attempt_id, "embedding", config["id"], config["version"], "running", db.now()))
    try:
        async with client_for(config) as client:
            response = await client.embeddings.create(model=config["model"], input=texts)
        ordered = sorted(response.data, key=lambda entry: entry.index)
        if len(ordered) != len(texts) or [entry.index for entry in ordered] != list(range(len(texts))):
            raise ValueError("embedding 响应数量或索引不匹配")
        vectors = [entry.embedding for entry in ordered]
        dimensions = len(vectors[0]) if vectors else 0
        if not dimensions or any(len(v) != dimensions or not all(math.isfinite(x) for x in v) for v in vectors):
            raise ValueError("embedding 维度或数值无效")
    except BaseException as exc:
        from .models_routes import safe_error
        with db.transaction() as conn:
            conn.execute("UPDATE model_calls SET status='failed',error=?,finished_at=? WHERE id=?", (safe_error(exc), db.now(), call_id))
        raise
    with db.transaction() as conn:
        conn.execute("UPDATE model_calls SET status='completed',usage_json=?,finished_at=? WHERE id=?", (response.usage.model_dump_json() if response.usage else None, db.now(), call_id))
    return vectors


async def index_source(payload: dict):
    source_id = payload["source_id"]
    source = source_record(source_id)
    if "source_revision" in payload and payload["source_revision"] != source["revision"]:
        raise ValueError("来源输入版本已变化，请重建新版本")
    try:
        config = config_at(payload["model_id"], payload["model_version"]) if "model_id" in payload else config_for("embedding")
    except HTTPException as exc:
        with db.transaction() as conn:
            conn.execute("UPDATE knowledge_sources SET status='failed',error=?,updated_at=? WHERE id=?", (str(exc.detail), db.now(), source_id))
        raise
    revision = source["revision"]
    version_id = db.uid()
    with db.transaction() as conn:
        conn.execute("UPDATE knowledge_sources SET status='processing',error=NULL,updated_at=? WHERE id=?", (db.now(), source_id))
        conn.execute("INSERT INTO knowledge_versions VALUES (?,?,?,?,?,?,?,?)", (version_id, source_id, revision, config["id"], config["version"], None, "building", db.now()))
    try:
        if source["content_id"]:
            with closing(db.connect()) as conn:
                content = conn.execute("SELECT published_json FROM contents WHERE id=?", (source["content_id"],)).fetchone()
            if not content or not content["published_json"]:
                raise ValueError("内容尚未发布")
            body = json.loads(content["published_json"])
            text = body["title"] + "\n\n" + body["markdown"]
        else:
            data = await asyncio.to_thread(storage.read, source["object_id"])
            # Parsing is delegated to a short-lived, bounded subprocess. It cannot hold the API event loop.
            import subprocess, sys
            completed = await asyncio.to_thread(subprocess.run,
                [sys.executable, "-m", "lyph.knowledge", source["filename"]], input=data,
                stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=60, check=False)
            if completed.returncode:
                raise ValueError(completed.stderr.decode("utf-8", errors="replace")[:300] or "文档解析失败")
            text = completed.stdout.decode("utf-8")
        chunks = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=150, separators=["\n\n", "\n", "。", " ", ""]).split_text(text)
        if len(chunks) > 3000:
            raise ValueError("切片超过 3000 个，请拆分来源")
        dimensions = None
        for start in range(0, len(chunks), 32):
            batch = chunks[start:start + 32]
            vectors = await embed_texts(config, batch, request_id=source_id)
            if dimensions is None:
                dimensions = len(vectors[0])
            if any(len(vector) != dimensions for vector in vectors):
                raise ValueError("embedding 在批次之间改变维度")
            with db.transaction() as conn:
                for offset, (chunk, vector) in enumerate(zip(batch, vectors, strict=True)):
                    conn.execute("INSERT INTO knowledge_chunks VALUES (?,?,?,?,?,?)", (db.uid(), source_id, version_id, start + offset, chunk, sqlite_vec.serialize_float32(vector)))
        with db.transaction() as conn:
            current = conn.execute("SELECT revision FROM knowledge_sources WHERE id=?", (source_id,)).fetchone()
            if not current or current["revision"] != revision:
                raise ValueError("来源已变更，候选索引未激活")
            conn.execute("UPDATE knowledge_versions SET status='ready',dimensions=? WHERE id=?", (dimensions, version_id))
            conn.execute("UPDATE knowledge_sources SET active_version=?,status='ready',text=?,error=NULL,updated_at=? WHERE id=?", (version_id, text, db.now(), source_id))
        return {"source_id": source_id, "version_id": version_id, "chunks": len(chunks), "dimensions": dimensions}
    except BaseException as exc:
        from .models_routes import safe_error
        message = str(exc)[:300] if isinstance(exc, ValueError) else safe_error(exc)
        with db.transaction() as conn:
            conn.execute("UPDATE knowledge_versions SET status='failed' WHERE id=?", (version_id,))
            conn.execute("UPDATE knowledge_sources SET status='failed',error=?,updated_at=? WHERE id=?", (message, db.now(), source_id))
        raise


async def retrieve(query: str, owner: bool = False, limit: int = 5, request_id: str | None = None, attempt_id: str | None = None):
    with closing(db.connect()) as conn:
        sources = [dict(row) for row in conn.execute("SELECT * FROM knowledge_sources WHERE active_version IS NOT NULL")]
    allowed = [source for source in sources if live_allowed(source, owner)]
    if not allowed:
        return []
    config_for("embedding")
    groups = {}
    with closing(db.connect()) as conn:
        for source in allowed:
            version = conn.execute("SELECT * FROM knowledge_versions WHERE id=?", (source["active_version"],)).fetchone()
            if version:
                groups.setdefault((version["model_id"], version["model_version"]), []).append(source["active_version"])
    rows = []
    for (model_id, model_version), versions in groups.items():
        # The old index remains usable after a failed rebuild, even across model changes.
        config = config_at(model_id, model_version)
        vector = (await embed_texts(config, [query], request_id, attempt_id))[0]
        with closing(db.connect()) as conn:
            conn.enable_load_extension(True)
            sqlite_vec.load(conn)
            conn.enable_load_extension(False)
            marks = ",".join("?" for _ in versions)
            rows.extend(conn.execute(f"""SELECT c.id,c.source_id,c.text,s.filename,s.content_id,
                vec_distance_cosine(c.embedding,?) AS distance FROM knowledge_chunks c
                JOIN knowledge_sources s ON s.id=c.source_id JOIN knowledge_versions v ON v.id=c.version_id
                WHERE c.version_id IN ({marks}) AND v.dimensions=?
                ORDER BY distance LIMIT ?""", (sqlite_vec.serialize_float32(vector), *versions, len(vector), limit)).fetchall())
    # Recheck after network time; a revocation during embedding must take effect immediately.
    safe_rows = []
    for row in rows:
        try:
            if live_allowed(source_record(row["source_id"]), owner):
                safe_rows.append(dict(row))
        except HTTPException:
            continue
    return sorted(safe_rows, key=lambda row: row["distance"])[:limit]


def enqueue_index(source_id: str, force: bool = False):
    source = source_record(source_id)
    config = config_for("embedding")
    key = f"knowledge:{source_id}:{source['revision']}:{config['id']}:{config['version']}"
    if force:
        key += ":" + db.uid()
    task = tasks.enqueue("knowledge.index", {"source_id": source_id, "source_revision": source["revision"],
                         "model_id": config["id"], "model_version": config["version"]}, key)
    with db.transaction() as conn:
        conn.execute("UPDATE knowledge_sources SET task_id=? WHERE id=?", (task["id"], source_id))
    return task


def sync_content(content_id: str):
    """Called by content publishing; errors remain visible without failing the publication."""
    with closing(db.connect()) as conn:
        content = conn.execute("SELECT * FROM contents WHERE id=?", (content_id,)).fetchone()
    if not content or not content["published_json"]:
        return None
    published = json.loads(content["published_json"])
    with db.transaction() as conn:
        old = conn.execute("SELECT * FROM knowledge_sources WHERE content_id=?", (content_id,)).fetchone()
        source_id = old["id"] if old else db.uid()
        revision = old["revision"] + 1 if old else 1
        conn.execute("""INSERT INTO knowledge_sources (id,content_id,filename,revision,status,created_at,updated_at)
            VALUES (?,?,?,?,?,?,?) ON CONFLICT(content_id) DO UPDATE SET filename=excluded.filename,
            revision=excluded.revision,status='pending',updated_at=excluded.updated_at""",
            (source_id, content_id, published["title"], revision, "pending", db.now(), db.now()))
    try:
        return enqueue_index(source_id)
    except HTTPException as exc:
        with db.transaction() as conn:
            conn.execute("UPDATE knowledge_sources SET status='failed',error=? WHERE id=?", (str(exc.detail), source_id))
        return None


@router.get("")
def list_sources():
    with closing(db.connect()) as conn:
        return {"items": [dict(row) for row in conn.execute("SELECT id,object_id,content_id,filename,visibility,rag_enabled,revision,active_version,status,error,task_id,created_at,updated_at FROM knowledge_sources ORDER BY created_at DESC")]}


@router.post("", status_code=201)
async def upload_source(file: UploadFile = File(...), visibility: str = Form("private"), rag_enabled: bool = Form(False)):
    if visibility not in ("public", "private") or Path(file.filename or "").suffix.lower() not in (".md", ".txt", ".pdf", ".docx"):
        raise HTTPException(422, "文件格式或可见性不支持")
    # Validate model readiness before persisting a half-configured pipeline.
    config_for("embedding")
    data = await file.read(settings.max_upload_bytes + 1)
    if len(data) > settings.max_upload_bytes:
        raise HTTPException(413, "文件超过 20MiB")
    if not data:
        raise HTTPException(422, "文件为空")
    media = await asyncio.to_thread(storage.put, data, file.filename, file.content_type or "application/octet-stream")
    source_id = db.uid()
    with db.transaction() as conn:
        filename = Path((file.filename or "source.txt").replace("\\", "/")).name[:240]
        conn.execute("INSERT INTO knowledge_sources (id,object_id,filename,visibility,rag_enabled,status,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?)", (source_id, media["id"], filename, visibility, int(rag_enabled), "pending", db.now(), db.now()))
        storage.add_reference(conn, media["id"], "knowledge", source_id, "original")
    task = enqueue_index(source_id)
    db.audit("knowledge.upload", source_id, {"object_id": media["id"]})
    return source_record(source_id) | {"task": task}


@router.get("/{source_id}")
def detail(source_id: str):
    return source_record(source_id)


@router.get("/{source_id}/chunks")
def chunks(source_id: str):
    source = source_record(source_id)
    with closing(db.connect()) as conn:
        return {"items": [dict(row) for row in conn.execute("SELECT id,position,text FROM knowledge_chunks WHERE version_id=? ORDER BY position", (source["active_version"],))]}


@router.post("/{source_id}/rebuild")
def rebuild(source_id: str):
    return enqueue_index(source_id, force=True)


class PermissionInput(BaseModel):
    visibility: str | None = None
    rag_enabled: bool


@router.put("/{source_id}/permissions")
def permissions(source_id: str, payload: PermissionInput):
    source = source_record(source_id)
    if source["content_id"]:
        raise HTTPException(409, "内容来源请在内容编辑页面设置权限")
    visibility = payload.visibility or source["visibility"]
    if visibility not in ("public", "private"):
        raise HTTPException(422, "可见性无效")
    with db.transaction() as conn:
        conn.execute("UPDATE knowledge_sources SET visibility=?,rag_enabled=?,updated_at=? WHERE id=?", (visibility, int(payload.rag_enabled), db.now(), source_id))
    db.audit("knowledge.permissions", source_id, payload.model_dump())
    return source_record(source_id)


@router.delete("/{source_id}")
def delete_source(source_id: str):
    source_record(source_id)
    with db.transaction() as conn:
        conn.execute("DELETE FROM knowledge_chunks WHERE source_id=?", (source_id,))
        conn.execute("DELETE FROM knowledge_versions WHERE source_id=?", (source_id,))
        conn.execute("DELETE FROM knowledge_sources WHERE id=?", (source_id,))
        conn.execute("DELETE FROM file_references WHERE owner_type='knowledge' AND owner_id=?", (source_id,))
    db.audit("knowledge.delete", source_id, {})
    return {"deleted": True}


class QueryInput(BaseModel):
    query: str = Field(min_length=1, max_length=4000)
    identity: str = "visitor"
    scope: str | None = None


@router.post("/search")
@router.post("/test/search")
async def search(payload: QueryInput):
    identity = payload.scope or payload.identity
    if identity not in ("owner", "visitor"):
        raise HTTPException(422, "检索身份无效")
    return {"items": await retrieve(payload.query, owner=identity == "owner")}


if __name__ == "__main__":
    import sys
    try:
        # Linux RLIMIT_AS/CPU bound extraction; Windows uses parent timeout.
        if sys.platform != "win32":
            import resource
            resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024, 512 * 1024 * 1024))
            resource.setrlimit(resource.RLIMIT_CPU, (40, 40))
        sys.stdout.buffer.write(extract_text(sys.stdin.buffer.read(), sys.argv[1]).encode("utf-8"))
    except Exception as exc:
        sys.stderr.buffer.write((str(exc)[:300] if isinstance(exc, ValueError) else type(exc).__name__).encode("utf-8"))
        sys.exit(1)
